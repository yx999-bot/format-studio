"""Structure checks for Format Studio output (python-docx + raw XML)."""

import glob
import json
import os
import sys
import zipfile

import docx
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, 'out')


def east_asia_font(run):
    rpr = run._element.find(
        '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rPr')
    if rpr is None:
        return None
    fonts = rpr.find(
        '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts')
    if fonts is None:
        return None
    return fonts.get(
        '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia')


def describe_paragraph(paragraph):
    pf = paragraph.paragraph_format
    run = paragraph.runs[0] if paragraph.runs else None
    line = pf.line_spacing
    return {
        'text': (paragraph.text or '')[:26],
        'font': east_asia_font(run) if run is not None else None,
        'size_pt': (run.font.size.pt if run is not None and run.font.size else None),
        'bold': (run.font.bold if run is not None else None),
        'align': str(pf.alignment) if pf.alignment is not None else None,
        'first_line_indent_pt': (pf.first_line_indent.pt if pf.first_line_indent else None),
        'line_spacing': (line.pt if hasattr(line, 'pt') else line),
        'space_before_pt': (pf.space_before.pt if pf.space_before else None),
        'space_after_pt': (pf.space_after.pt if pf.space_after else None),
    }


def check(path):
    report = {'file': os.path.basename(path)}
    try:
        document = docx.Document(path)
    except Exception as exc:  # noqa: BLE001
        report['error'] = f'python-docx could not open: {exc}'
        return report

    section = document.sections[0]
    report['page'] = {
        'width_mm': round(section.page_width.mm, 2),
        'height_mm': round(section.page_height.mm, 2),
        'top_mm': round(section.top_margin.mm, 2),
        'bottom_mm': round(section.bottom_margin.mm, 2),
        'left_mm': round(section.left_margin.mm, 2),
        'right_mm': round(section.right_margin.mm, 2),
    }
    report['paragraph_count'] = len(document.paragraphs)
    report['paragraphs'] = [describe_paragraph(p) for p in document.paragraphs[:10]]

    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        report['parts'] = names
        document_xml = archive.read('word/document.xml').decode('utf-8')
        report['footer_reference'] = 'footerReference' in document_xml
        report['has_page_field'] = 'PAGE' in (archive.read('word/footer1.xml').decode('utf-8')
                                              if 'word/footer1.xml' in names else '')
        report['first_line_chars'] = 'firstLineChars="200"' in document_xml
        report['east_asia_headers'] = document_xml.count('w:eastAsia=')
    return report


def main():
    paths = sorted(glob.glob(os.path.join(OUT, '*.docx')))
    if not paths:
        print('no docx files found - run `node test/validate.mjs` first')
        return 1
    reports = [check(path) for path in paths]
    print(json.dumps(reports, ensure_ascii=False, indent=2))
    failures = [r for r in reports if 'error' in r]
    return 1 if failures else 0


if __name__ == '__main__':
    sys.exit(main())
